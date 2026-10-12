#!/usr/bin/env python3
"""協会けんぽ「保険料額表」Excel（全都道府県分）から社会保険料データ JS を生成する。

使い方:
  python convert_si.py --xlsx source/kenpo/r8ippan3.xlsx --fiscal-year 令和8年度 --year 2026 --out tables
雇用保険料率などの定数はこのスクリプト内（EMPLOYMENT）で管理する。
"""
import argparse, datetime, json, re
from pathlib import Path
import openpyxl

# 雇用保険料率（労働者負担・%）令和8年度（2026-04〜2027-03）厚生労働省リーフレット
EMPLOYMENT = {
    "fiscalYear": "令和8年度", "from": "2026-04",
    "options": [["一般の事業", 0.5], ["農林水産・清酒製造の事業", 0.6], ["建設の事業", 0.6]],
    "source": "https://www.mhlw.go.jp/stf/seisakunitsuite/bunya/0000108634.html",
}

def pref_name(sheet):
    if sheet == "北海道": return sheet
    if sheet == "東京": return "東京都"
    if sheet in ("大阪", "京都"): return sheet + "府"
    return sheet + "県"

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", type=Path, required=True)
    ap.add_argument("--fiscal-year", required=True)
    ap.add_argument("--year", type=int, required=True)
    ap.add_argument("--source-url", default="https://www.kyoukaikenpo.or.jp/assets/r8ippan3.xlsx")
    ap.add_argument("--out", type=Path, default=Path("tables"))
    a = ap.parse_args()
    wb = openpyxl.load_workbook(str(a.xlsx), data_only=True)
    prefs, grades_ref, care, support, pension = [], None, None, None, None
    for ws in wb.worksheets:
        rate = round(float(ws.cell(8, 6).value) * 100, 2)
        rate_care = round(float(ws.cell(8, 8).value) * 100, 2)
        sup = round(float(ws.cell(8, 10).value) * 100, 2)
        pen = round(float(ws.cell(8, 12).value) * 100, 2)
        c = round(rate_care - rate, 2)
        for name, old, new in (("介護", care, c), ("支援金", support, sup), ("厚年", pension, pen)):
            assert old is None or old == new, f"{ws.title}: {name}率が他県と異なる {old} vs {new}"
        care, support, pension = c, sup, pen
        g = []
        for r in range(11, 61):
            label, std, lo, hi = (ws.cell(r, k).value for k in (1, 2, 3, 5))
            m = re.match(r"(\d+)(?:\((\d+)\))?", str(label))
            g.append([int(m.group(1)), int(std), int(lo) if lo else 0, int(hi) if hi else None,
                      int(m.group(2)) if m.group(2) else None])
        assert len(g) == 50 and g[0][0] == 1 and g[-1][0] == 50, f"{ws.title}: 等級が50行でない"
        for x, y in zip(g, g[1:]):
            assert x[3] == y[2], f"{ws.title}: 報酬月額の範囲が連続していない {x} {y}"
        if grades_ref is None:
            grades_ref = g
        assert g == grades_ref, f"{ws.title}: 等級表が他県と異なる"
        prefs.append([pref_name(ws.title), rate])
    assert len(prefs) == 47
    data = {
        "year": a.year, "fiscalYear": a.fiscal_year, "appliesFrom": f"{a.year}-03",
        "retrieved": datetime.date.today().isoformat(), "source": a.source_url,
        "kenpo": {"prefectures": prefs, "careRate": care, "supportRate": support,
                  "supportFrom": f"{a.year}-04", "grades": grades_ref, "bonusCapYearly": 5730000},
        "pension": {"rate": pension, "bonusCapMonthly": 1500000,
                    "minGradeUpper": 93000, "maxGradeLower": 635000},
        "employment": EMPLOYMENT,
    }
    a.out.mkdir(parents=True, exist_ok=True)
    js = json.dumps(data, ensure_ascii=False, separators=(",", ":"))
    (a.out / f"si-{a.year}.json").write_text(json.dumps(data, ensure_ascii=False) + "\n", encoding="utf-8")
    (a.out / f"si-{a.year}.js").write_text(
        f"// 自動生成: tools/convert_si.py（{data['retrieved']}）。手で編集しない。\n"
        f"window.GENSEN_SI = window.GENSEN_SI || {{}};\nwindow.GENSEN_SI[\"{a.year}\"] = {js};\n", encoding="utf-8")
    print(f"{a.fiscal_year}: 都道府県 {len(prefs)} 件（東京都 {dict(prefs)['東京都']}%）, 介護 {care}%, 支援金 {support}%, 厚年 {pension}%, 等級 {len(grades_ref)}")

if __name__ == "__main__":
    main()
